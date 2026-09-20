"""Export paper-main CSV outputs to an Excel workbook and 118-bus table."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


ROOT = Path(__file__).resolve().parent
PAPER_DIR = ROOT.parent
OUT_DIR = ROOT / "results" / "paper_main"
FIG_DIR = PAPER_DIR / "fig"


CSV_SHEETS = {
    "30_benchmark": OUT_DIR / "benchmark_summary_30base.csv",
    "30_unit_loc": OUT_DIR / "unit_level_profit_loc_30base.csv",
    "30_ramping": OUT_DIR / "ramping_diagnostic_30base.csv",
    "30_convergence": OUT_DIR / "convergence_profiles_30base.csv",
    "30_price_output": OUT_DIR / "prices_schedule_all_units_30base.csv",
    "30_schedule": OUT_DIR / "schedule_summary_30base.csv",
    "30_lines": OUT_DIR / "line_utilization_audit_30base.csv",
    "30_generators": OUT_DIR / "generator_params_30base.csv",
    "118_benchmark_fast": OUT_DIR / "benchmark_118_fast.csv",
}


METHOD_LABEL = {
    "lmp": "LMP",
    "mirp": "IRP",
    "level": "LVM",
    "dwp": "DWP",
    "dwp_incremental": "DWP-inc",
    "xiao": "S-CHP",
    "chp": "D-CHP",
}

ORDER = ["lmp", "mirp", "level", "dwp", "dwp_incremental", "xiao", "chp"]
ORDER_118 = ["lmp", "mirp", "level", "dwp_incremental", "xiao", "chp"]


def read_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def fnum(value, default=""):
    try:
        if value in ("", None):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def fmt_money(value) -> str:
    number = fnum(value, None)
    return "--" if number is None else f"{number:,.2f}"


def fmt_time(value) -> str:
    number = fnum(value, None)
    return "--" if number is None else f"{number:,.2f}"


def fmt_iter(value) -> str:
    number = fnum(value, None)
    return "--" if number is None else f"{int(round(number)):,}"


def write_workbook() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, path in CSV_SHEETS.items():
        ws = wb.create_sheet(sheet_name[:31])
        rows = read_csv(path)
        if not rows:
            ws.append([f"Missing or empty: {path}"])
            continue
        headers = list(rows[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row.get(h, "") for h in headers])
        for col in ws.columns:
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 42)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_DIR / "paper_main_results.xlsx")


def write_118_table() -> None:
    rows = read_csv(OUT_DIR / "benchmark_118_fast.csv")
    by_method = {row["method"]: row for row in rows}
    lines = [
        r"\begin{table*}[t]",
        r"\centering",
        r"\caption{Scalability results on the IEEE 118-bus UC case with PTDF constraints.}",
        r"\label{tab:scalability-118}",
        r"\scriptsize",
        r"\begin{tabular}{lrrrrrrrr}",
        r"\toprule",
        (
            r"Method & $C^{\mathrm{LP}}$ (\$) & Gen. LOC (\$) & FTR (\$) & "
            r"Total uplift (\$) & Time (s) & Build (s) & Solver (s) & Iter. \\"
        ),
        r"\midrule",
    ]
    for method in ORDER_118:
        row = by_method.get(method)
        if not row:
            continue
        lines.append(
            " & ".join(
                [
                    METHOD_LABEL.get(method, method),
                    fmt_money(row.get("pricing_obj")),
                    fmt_money(row.get("gen_uplift")),
                    fmt_money(row.get("ftr_cost")),
                    fmt_money(row.get("total_uplift")),
                    fmt_time(row.get("method_time")),
                    fmt_time(row.get("build_time")),
                    fmt_time(row.get("solver_time")),
                    fmt_iter(row.get("n_iter")),
                ]
            )
            + r" \\"
        )
    lines.extend([r"\bottomrule", r"\end{tabular}", r"\end{table*}", ""])
    text = "\n".join(lines)
    (OUT_DIR / "table_118_scalability.tex").write_text(text, encoding="utf-8")
    (FIG_DIR / "table_118_scalability.tex").write_text(text, encoding="utf-8")


def main() -> None:
    write_workbook()
    write_118_table()
    print(f"Wrote {OUT_DIR / 'paper_main_results.xlsx'}")
    print(f"Wrote {FIG_DIR / 'table_118_scalability.tex'}")


if __name__ == "__main__":
    main()
